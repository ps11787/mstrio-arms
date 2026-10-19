from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUTPUT_PATH = r"C:\Users\psokulski\Documents\Projects\BFA\mstrio-arms-kroschu\docs\BfA_Import_Package_Script.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    if level == 1:
        heading = doc.add_heading(text, level=1)
    elif level == 2:
        heading = doc.add_heading(text, level=2)
    else:
        heading = doc.add_heading(text, level=3)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_paragraph(doc: Document, text: str, bold: bool = False, font_size: int = 11) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run(text)
    if not paragraph.runs:
        run = paragraph.add_run(text)
    if bold:
        run.bold = True
    run.font.size = Pt(font_size)
    if not paragraph.runs:
        paragraph.add_run(text)


def add_label_value(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    run_label = paragraph.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.size = Pt(11)
    run_value = paragraph.add_run(value)
    run_value.font.size = Pt(11)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for idx, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        if idx < len(lines) - 1:
            run = paragraph.add_run("\n")
            run.font.name = "Consolas"
            run.font.size = Pt(10)


def build_document() -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    doc.core_properties.title = "BfA Import Package Script"
    doc.core_properties.subject = "Technical documentation for MicroStrategy migration import automation"
    doc.core_properties.author = "BfA"
    doc.core_properties.company = "BfA"
    doc.core_properties.keywords = "MicroStrategy; migration; ARMS; import script"

    # Referenzstruktur zur Vorlage aus dem Referenz-Dokument
    doc.add_paragraph("Creative Work – - /")
    doc.add_paragraph("Name of the work")
    doc.add_paragraph("Name of object: Import Package Script")
    doc.add_paragraph("Version of Work: 1.0")
    doc.add_paragraph("Project: BfA / MicroStrategy ARMS")
    doc.add_paragraph("Type of the work")
    doc.add_paragraph("Description of the work")

    description = (
        "Imports an existing MicroStrategy migration package into the target project. "
        "The script resolves the migration object GUID from the ARMS request metadata, "
        "opens the target environment connection, executes the import, polls the migration status, "
        "updates the request and migration tables, and logs the result for operational traceability."
    )
    doc.add_paragraph(description)

    doc.add_paragraph("Co-authors")
    doc.add_paragraph("If applicable, else N/A")
    doc.add_paragraph("Repository")
    doc.add_paragraph("mstrio-arms-kroschu / import_package.py")

    add_heading(doc, "1. Purpose", 1)
    doc.add_paragraph(
        "The script is the downstream automation step after package creation. It reads the active request row, "
        "resolves the migration object GUID, identifies the target project and environment, imports the package into the target project, "
        "and updates the ARMS request status to reflect the end state."
    )

    add_heading(doc, "2. Scope and responsibilities", 1)
    doc.add_paragraph(
        "This script consumes already-generated migration metadata and does not create the package itself. It assumes the earlier workflow has already created a valid migration object and stored the request context in the ARMS tables."
    )
    doc.add_paragraph("Main responsibilities:")
    bullets = [
        "Read the current migration request from the request table.",
        "Resolve the migration object GUID and source project name.",
        "Determine the correct target project name and environment.",
        "Establish a connection to the target MicroStrategy environment.",
        "Execute the package import with Migration.migrate().",
        "Poll the migration status until it reaches a terminal state.",
        "Write the target status to the ARMS request and migration tables.",
        "Insert a comment for auditability in case of success or failure.",
    ]
    for item in bullets:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "3. Configuration constants", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    header_cells[0].text = "Configuration item"
    header_cells[1].text = "Value / purpose"

    rows = [
        ("BASE_URL", "MicroStrategy Library URL used as the source environment base URL."),
        ("SOURCE_ENVIRONMENT", "Current source environment name, default value: dev."),
        ("ENVIRONMENT_URLS", "Map of environment names to their target MicroStrategy URLs."),
        ("MSTR_USERNAME", "MicroStrategy user name used for both source and target connections."),
        ("MSTR_PASSWORD", "Credential used for the connection objects."),
        ("PROJECT_ID", "MicroStrategy project identifier used for ARMS datasource calls."),
        ("REQUEST_REPORT_ID", "Reference to the report used to read a request row."),
        ("DATASOURCE_INSTANCE_ID", "Datasource instance used to query the ARMS metadata tables."),
        ("TARGET_PROJECT_TABLE", "Lookup table containing target project names and status mapping."),
        ("STATUS_TABLE", "Lookup table for status metadata and environment mapping."),
        ("TARGET_ENV_TABLE", "Lookup table for target environment details such as prefix and URL."),
        ("REQUEST_TABLE", "ARMS request table used for the current migration workflow."),
        ("MIGRATION_OBJECT_TABLE", "Lookup table used to resolve the migration object GUID."),
        ("REQUEST_STATUS_FILTER", "Statuses considered active for processing: 5, 7, 9."),
    ]
    for key, value in rows:
        row = table.add_row().cells
        row[0].text = key
        row[1].text = value

    add_heading(doc, "4. Execution flow", 1)
    steps = [
        "Create a source connection with Connection(...).",
        "Open the datasource instance for ARMS queries.",
        "Select the first request row whose status matches REQUEST_STATUS_FILTER.",
        "Resolve the migration object GUID from the request row and migration lookup table.",
        "Read the source project name, target project name, and request status data.",
        "Determine the target environment and target project name using lookup tables and naming rules.",
        "Create a target connection to the resolved target environment project.",
        "Instantiate Migration(my_conn, id=migration_obj_guid).",
        "Call current_migration.migrate(target_env=target_env, target_project_name=target_project_name, generate_undo=True).",
        "Poll until the import reaches a terminal state via _wait_for_migration_completion().",
        "Insert the migration record into the ARMS dimension table.",
        "Write a success comment and update the request status using status_id_success from the target environment.",
        "Close the target connection in the finalization block.",
    ]
    for idx, step in enumerate(steps, start=1):
        doc.add_paragraph(f"{idx}. {step}")

    add_heading(doc, "5. Main helper functions", 1)
    add_label_value(doc, "_normalize_lookup_key", "Normalizes field names to compare query response keys regardless of casing and spacing differences.")
    add_label_value(doc, "_get_first_value_from_result_data", "Returns the first non-empty value from a result dataset, supporting list-based and scalar responses.")
    add_label_value(doc, "_rows_from_result", "Converts raw datasource output into a standard list of dictionaries for downstream processing.")
    add_label_value(doc, "_resolve_migration_object_guid", "Looks up the migration object GUID using the request migration_key_id and the migration lookup table.")
    add_label_value(doc, "_get_target_environment_details", "Loads target environment metadata such as name, prefix, and URL from the environment lookup table.")
    add_label_value(doc, "_derive_target_project_name", "Builds the target project name from the prefix and source project name, while removing redundant prefixes.")
    add_label_value(doc, "_get_result_status_id_for_environment", "Reads and validates status_id_success or status_id_failed from the selected target environment.")
    add_label_value(doc, "_insert_migration_dim_record", "Appends a successful or failed migration attempt without overwriting previous history.")
    add_label_value(doc, "_insert_comment", "Adds an audit comment to the ARMS comments table.")

    add_heading(doc, "6. Key data logic", 1)
    doc.add_paragraph(
        "The script relies on a layered lookup model: request metadata → migration lookup table → target environment table → target project table → status table. "
        "This makes the import workflow deterministic and allows environment-specific naming and status mapping without hardcoded project names in every step."
    )

    add_heading(doc, "7. Validation and failure handling", 1)
    failures = [
        "No active request rows are returned: the script raises ValueError and stops.",
        "No migration object GUID is found: the script aborts before the migration starts.",
        "The target project name is missing: the script cannot create a valid target connection.",
        "Missing or invalid status_id_success/status_id_failed on the target environment stops import before it starts.",
        "If import fails, the request is updated to status_id_failed, a failed history record and error comment are inserted, and the exception is re-raised.",
        "A locked migration triggers package reuse into a new target-bound migration. Its ID is persisted before import starts; its binding is checked before submission.",
        "Every SQL query is printed immediately before execution.",
        "If the import succeeds but the request status update fails, the import may already exist while the request table remains out of sync.",
    ]
    for item in failures:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "8. Representative code", 1)
    code = [
        "request_rows = _fetch_requests_for_statuses(db_instance, PROJECT_ID, REQUEST_STATUS_FILTER)",
        "migration_obj_guid = _resolve_migration_object_guid(db_instance, PROJECT_ID, request_row) or \"\"",
        "target_env_details = _get_target_environment_details(db_instance, target_env_id, PROJECT_ID)",
        "import_result = current_migration.migrate(",
        "    target_env=target_env,",
        "    target_project_name=target_project_name,",
        "    generate_undo=True,",
        ")",
        "_wait_for_migration_completion(current_migration)",
        "db_instance.execute_query(",
        "    query=f\"UPDATE arms.t_mig_dim_request SET status_id={target_status_id} WHERE request_id={request_id};\",",
        "    project_id=PROJECT_ID,",
        ")",
    ]
    add_code_block(doc, code)

    add_heading(doc, "9. Operational notes", 1)
    notes = [
        "The script expects that the ARMS lookup tables are already populated and consistent with the target project naming model.",
        "The script reads only the first matching request row and treats it as the active migration context.",
        "Target URLs are resolved from a centralized map; if no explicit environment mapping is defined, the script falls back to the source base URL.",
        "Connections are always closed in the finalization block, even if the migration or status update fails.",
        "Credentials are embedded in the script and should be managed according to the environment's security policy.",
    ]
    for item in notes:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()
    add_heading(doc, "Appendix: script path", 1)
    doc.add_paragraph("File: import_package.py")
    doc.add_paragraph("Repository: C:/Users/psokulski/Documents/Projects/BFA/mstrio-arms-kroschu")
    doc.add_paragraph("Related documentation: docs/import_package.md")

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT_PATH)
    print(f"Word document created: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
